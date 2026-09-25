"""Generate synthetic renderer fixtures into a new directory (optional developer tool)."""
import argparse
from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.pdfgen import canvas
import xlsxwriter


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    args.destination.mkdir(parents=True, exist_ok=False)
    target = args.destination

    pdf = canvas.Canvas(str(target / "review.pdf"))
    for text in ["Board summary", "Finance and HR review required", "Cash runway: 18 months"]:
        pdf.drawString(60, 700, text)
        pdf.showPage()
    pdf.save()

    deck = Presentation()
    for title, body in [
        ("Mesh Alpha Review", "Exact native review"),
        ("People and finance decision", "Approve three hires after the revised cash-flow review."),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    deck.save(target / "review.pptx")

    document = Document()
    document.add_heading("People plan", 0)
    document.add_paragraph("Review staffing before approval.")
    document.save(target / "review.docx")

    with xlsxwriter.Workbook(target / "review.xlsx") as workbook:
        sheet = workbook.add_worksheet("Budget")
        sheet.write_row("A1", ["Category", "Revenue", "Cost", "Balance"])
        sheet.write_row("A2", ["Revenue", 120000, 85000])
        sheet.write_formula("D2", "=B2-C2", None, 35000)


if __name__ == "__main__":
    main()
